import os
from dotenv import load_dotenv
load_dotenv()

import re
import json
import base64
import time
import openpyxl
from anthropic import Anthropic
from pdf2image import convert_from_path
from pathlib import Path
from io import BytesIO

# 설정
DATA_DIR      = os.environ["DATA_DIR"]
CODE_DIR      = os.path.dirname(os.path.abspath(__file__))

ANTHROPIC_API_KEY = os.environ["ANTHROPIC_API_KEY"]
TEST_FOLDER   = os.path.join(DATA_DIR, "test_folder")
OUTPUT_EXCEL  = os.path.join(DATA_DIR, "테스트세트.xlsx")
TAXONOMY_PATH = os.path.join(CODE_DIR, "taxonomy.md")

MODEL = "claude-haiku-4-5-20251001"
EXAM_TYPES = ["고졸", "중졸", "초졸"]
SUBJECT_ORDER = ["국어", "수학", "영어", "사회", "과학"]
BATCH_SIZE = 15

PILOT_TARGET = None
PILOT_PAGE_LIMIT = None

client = Anthropic(api_key=ANTHROPIC_API_KEY)

# 과목별 판단 기준
SUBJECT_NOTES = {
    "과학": """
[분류 시 참고할 판단 기준]
- 화살표(→) 양쪽이 같은 물질(상태만 다름)이면 "물질의 상태 변화", 다른 물질(새 물질 생성)이면 "화학 반응의 규칙과 에너지 변화"(중등) 또는 해당 반응 카테고리(고등)
- 세포호흡 반응식(포도당+산소→이산화탄소+물+에너지)은 화학이 아니라 생물로 분류 (물질대사/광합성과 호흡 등)
- "공유결합/이온결합" 개념 자체를 묻는 문제는 화학 > 화학 결합 (원자의 구조 아님)
- "물질대사"라는 용어 자체를 정의하는 문제만 생물 > 물질대사. 광합성/호흡의 구체적 과정은 기존 카테고리(광합성과 호흡) 유지
- 생태계, 생물다양성, 환경 요인에 따른 생물의 적응/진화(자연선택 포함)는 대분류가 "지구과학"이고 중분류는 "생태계" 또는 "다양성"
- 연료전지 문제: "무슨 반응/기체가 나오는가" 물으면 산화 환원, "무슨 발전 방식인가" 물으면 재생에너지(발전방식)
- "태양 고도"나 "낮의 길이" 개념이 나오면 계절의 변화. 하루 단위 자전 현상(낮밤, 동→서로 움직이는 것처럼 보임)은 지구와 달. 다른 행성/별 비교는 태양계와 별
- 힘의 종류(중력/마찰력/탄성력 등) 자체를 묻는 문제는 여러 가지 힘. 속력/위치·운동에너지 계산은 운동과 에너지
- 위치·운동에너지만 다루면 운동과 에너지(고등은 역학적에너지), 전기/화학/열/빛 등 다른 종류 에너지 간 전환이면 에너지 전환과 보존
- 광물의 특성(조흔색, 굳기, 염산 반응)·암석·지층·화산·지진은 지권의 변화. 바닷물·염류·해류·수온약층은 수권과 해수의 순환
""",
    "영어": """
[분류 시 참고할 판단 기준]
- 지문(대화/글) 안의 사실 정보를 정확히 찾아 확인/대조하는 문제 → 독해 > 세부 내용 파악
- 빈칸에 들어갈 자연스러운 대화 표현(질문이든 대답이든)을 고르는 문제, 실생활 상황 대화 → 생활영어
- 지문 없이 그림+단어만으로 판단하는 문제(그림-낱말 매칭, 공통 철자 찾기 등) → 어휘 > 단어
- 두 단어 사이의 의미 관계(유의어/반의어/상위-하위어 등) 유추 → 어휘 > 두 단어의 관계
- 비교급/최상급 문법 형태를 채우거나 판단 → 문법 > 형용사, 부사, 비교
- 대명사·지시어가 가리키는 대상을 찾는 문제 → 독해 > 지칭 추론
- 광고문/안내문/초대장/일정표 등 실물 문서 형식 지문(대화 형식 아님) → 독해 > 실용문
""",
}

# taxonomy 파싱
def parse_taxonomy(path: str) -> dict:
    with open(path, "r", encoding="utf-8") as f:
        text = f.read()

    result = {}
    section = None
    for line in text.split("\n"):
        m = re.match(r"^## (.+)", line)
        if m:
            section = m.group(1).strip()
            result[section] = {}
            continue
        if section and line.strip().startswith("- "):
            parts = line.strip()[2:].split(" > ")
            if len(parts) >= 2:
                large, mid = parts[0].strip(), parts[1].strip()
                result[section].setdefault(large, [])
                if mid not in result[section][large]:
                    result[section][large].append(mid)
            elif parts and parts[0].strip():
                large = parts[0].strip()
                result[section].setdefault(large, [])
    return result

TAXONOMY = parse_taxonomy(TAXONOMY_PATH)

def get_section_key(exam_type: str, subject: str) -> str:
    level_map = {"초졸": "초등", "중졸": "중등", "고졸": "고등"}
    level = level_map.get(exam_type, "")
    if subject == "영어":
        for key in TAXONOMY:
            if "영어" in key:
                return key
    return f"{level}_{subject}"

def chunk(lst, n):
    for i in range(0, len(lst), n):
        yield lst[i:i + n]

def image_to_base64(img):
    buffer = BytesIO()
    img.save(buffer, format="PNG")
    return base64.b64encode(buffer.getvalue()).decode("utf-8")


# 1단계: 이미지에서 문제/보기 추출
def extract_page(page_image, info, page_num, max_retry=3):
    tool = {
        "name": "extract_questions",
        "description": "페이지 안의 각 문제를 추출한다 (분류는 하지 않음).",
        "input_schema": {
            "type": "object",
            "properties": {
                "questions": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "question_number": {"type": "integer"},
                            "question": {"type": "string"},
                            "choices": {"type": "string"},
                            "has_image": {"type": "boolean"},
                        },
                        "required": ["question_number", "question", "choices", "has_image"],
                    },
                }
            },
            "required": ["questions"],
        },
    }

    prompt = f"""아래는 검정고시 {info['exam_type']} {info['subject']} 시험지 {page_num}페이지 이미지입니다.
이 페이지에 있는 각 문제의 문제 텍스트(보기 제외), 보기, 이미지 유무를 추출하세요.
- has_image는 문제에 도형/표/그래프 등 이미지가 있으면 true
- 보기가 이미지(도형/그래프 등)인 경우 choices는 빈 문자열

반드시 extract_questions 도구를 호출하세요."""

    b64 = image_to_base64(page_image)

    for attempt in range(max_retry):
        try:
            resp = client.messages.create(
                model=MODEL,
                max_tokens=4000,
                tools=[tool],
                tool_choice={"type": "tool", "name": "extract_questions"},
                messages=[{
                    "role": "user",
                    "content": [
                        {"type": "image", "source": {"type": "base64", "media_type": "image/png", "data": b64}},
                        {"type": "text", "text": prompt}
                    ]
                }],
            )
            for block in resp.content:
                if block.type == "tool_use":
                    return block.input["questions"]
        except Exception as e:
            print(f"  [재시도 {attempt+1}/{max_retry}] {e}")
            time.sleep(10)

    return []


# 2단계: 대분류 배치 분류
def classify_large_batch(items, large_list, subject_note, max_retry=3):
    tool = {
        "name": "assign_large_category",
        "description": "각 문제(row_key)에 대분류를 하나씩 배정한다.",
        "input_schema": {
            "type": "object",
            "properties": {
                "results": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "row_key": {"type": "integer"},
                            "category_large": {"type": "string", "enum": large_list},
                        },
                        "required": ["row_key", "category_large"],
                    },
                }
            },
            "required": ["results"],
        },
    }

    items_text = "\n\n".join(
        f"[row_key={it['row_key']}]\n문제: {it['question']}\n보기: {it['choices']}"
        for it in items
    )
    prompt = f"""아래 검정고시 문제들 각각에 대해 가장 적합한 대분류를 선택하세요.

사용 가능한 대분류 목록 (이 중에서만 선택):
{large_list}
{subject_note}
문제 목록:
{items_text}

각 문제(row_key)에 대해 category_large를 하나씩 배정하세요. 반드시 assign_large_category 도구를 호출하세요."""

    for attempt in range(max_retry):
        try:
            resp = client.messages.create(
                model=MODEL,
                max_tokens=4000,
                tools=[tool],
                tool_choice={"type": "tool", "name": "assign_large_category"},
                messages=[{"role": "user", "content": prompt}],
            )
            for block in resp.content:
                if block.type == "tool_use":
                    data = block.input
                    return {r["row_key"]: r["category_large"] for r in data["results"]}
        except Exception as e:
            print(f"  [대분류 재시도 {attempt + 1}/{max_retry}] {e}")
            time.sleep(8)
    return {}


# 3단계: 중분류 배치 분류
def classify_mid_batch(items, mid_list, subject_note, max_retry=3):
    tool = {
        "name": "assign_mid_category",
        "description": "각 문제(row_key)에 중분류를 하나씩 배정한다.",
        "input_schema": {
            "type": "object",
            "properties": {
                "results": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "row_key": {"type": "integer"},
                            "category_mid": {"type": "string", "enum": mid_list},
                        },
                        "required": ["row_key", "category_mid"],
                    },
                }
            },
            "required": ["results"],
        },
    }

    items_text = "\n\n".join(
        f"[row_key={it['row_key']}]\n문제: {it['question']}\n보기: {it['choices']}"
        for it in items
    )
    prompt = f"""아래 검정고시 문제들은 이미 대분류가 확정된 상태입니다.
각 문제에 대해 그 대분류에 속하는 중분류 중 가장 적합한 것을 선택하세요.

사용 가능한 중분류 목록 (이 중에서만 선택, 반드시 이 목록 안에서):
{mid_list}
{subject_note}
문제 목록:
{items_text}

각 문제(row_key)에 대해 category_mid를 하나씩 배정하세요. 반드시 assign_mid_category 도구를 호출하세요."""

    for attempt in range(max_retry):
        try:
            resp = client.messages.create(
                model=MODEL,
                max_tokens=4000,
                tools=[tool],
                tool_choice={"type": "tool", "name": "assign_mid_category"},
                messages=[{"role": "user", "content": prompt}],
            )
            for block in resp.content:
                if block.type == "tool_use":
                    data = block.input
                    return {r["row_key"]: r["category_mid"] for r in data["results"]}
        except Exception as e:
            print(f"  [중분류 재시도 {attempt + 1}/{max_retry}] {e}")
            time.sleep(8)
    return {}


# 정답표 파싱
def parse_all_answers(exam_type):
    answer_path = os.path.join(TEST_FOLDER, f"2026년도_2차시험_{exam_type}학력_정답.pdf")
    if not os.path.exists(answer_path):
        print(f"[답안없음] {answer_path}")
        return {}

    pages = convert_from_path(answer_path)
    all_answers = {}

    tool = {
        "name": "extract_answers",
        "description": "정답표에서 문항번호별 정답을 추출한다.",
        "input_schema": {
            "type": "object",
            "properties": {
                "answers": {"type": "object", "additionalProperties": {"type": "string"}}
            },
            "required": ["answers"],
        },
    }

    for idx, page_img in enumerate(pages):
        if idx >= len(SUBJECT_ORDER):
            break
        subject = SUBJECT_ORDER[idx]
        b64 = image_to_base64(page_img)

        prompt = """이 정답표 이미지에서 문항번호와 정답을 추출하세요.
정답은 원문자(①②③④)를 숫자(1,2,3,4)로 변환하세요.
반드시 extract_answers 도구를 호출하세요."""

        for attempt in range(3):
            try:
                resp = client.messages.create(
                    model=MODEL,
                    max_tokens=1000,
                    tools=[tool],
                    tool_choice={"type": "tool", "name": "extract_answers"},
                    messages=[{
                        "role": "user",
                        "content": [
                            {"type": "image", "source": {"type": "base64", "media_type": "image/png", "data": b64}},
                            {"type": "text", "text": prompt}
                        ]
                    }],
                )
                for block in resp.content:
                    if block.type == "tool_use":
                        all_answers[subject] = block.input["answers"]
                time.sleep(2)
                break
            except Exception as e:
                print(f"  [답안 재시도] {e}")
                time.sleep(10)

    return all_answers


# 엑셀 초기화
if Path(OUTPUT_EXCEL).exists():
    wb = openpyxl.load_workbook(OUTPUT_EXCEL)
    ws = wb.active
else:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "id", "exam_type", "year", "exam_round", "subject",
        "question_number", "question", "choices", "answer",
        "has_image", "image_path",
        "category_large", "category_mid", "category_small"
    ])

# 기존 id → 행 번호 매핑
id_to_row = {}
for row in ws.iter_rows(min_row=2):
    if row[0].value:
        id_to_row[row[0].value] = row[0].row

failed_log = []

if PILOT_TARGET:
    targets = [PILOT_TARGET]
    print(f"[파일럿 모드] {PILOT_TARGET} 만 처리, 페이지 최대 {PILOT_PAGE_LIMIT}장")
else:
    targets = [(et, subj) for et in EXAM_TYPES for subj in SUBJECT_ORDER]

answers_cache = {}

for exam_type, subject in targets:
    if exam_type not in answers_cache:
        print(f"\n=== {exam_type} 답안 로드 ===")
        answers_cache[exam_type] = parse_all_answers(exam_type)
    answers_by_subject = answers_cache[exam_type]

    pdf_path = os.path.join(TEST_FOLDER, f"2026년도_2차시험_{exam_type}학력_{subject}_문제.pdf")
    if not os.path.exists(pdf_path):
        print(f"[파일없음] {pdf_path}")
        continue

    check_id = f"{exam_type}_2026_2회_{subject}_1"
    if check_id in id_to_row and not PILOT_TARGET:
        print(f"[이미처리됨] {exam_type} {subject}")
        continue

    print(f"\n처리중: {exam_type} {subject}")
    pages = convert_from_path(pdf_path)
    if PILOT_TARGET:
        pages = pages[:PILOT_PAGE_LIMIT]

    answers = answers_by_subject.get(subject, {})
    print(f"  답안 매칭: {len(answers)}개")

    info = {"exam_type": exam_type, "subject": subject}

    # 1단계: 문제/보기 추출
    extracted = []
    row_key_counter = 0
    for page_num, page_img in enumerate(pages, start=1):
        print(f"  페이지 {page_num}/{len(pages)} 추출 중...")
        results = extract_page(page_img, info, page_num)
        if not results:
            print(f"  [스킵] 페이지 {page_num} 추출 실패")
            failed_log.append({
                "exam_type": exam_type, "subject": subject,
                "page_num": page_num, "사유": "페이지 추출 실패(응답 없음)"
            })
            continue
        time.sleep(3)
        for r in results:
            row_key_counter += 1
            extracted.append({
                "row_key": row_key_counter,
                "question_number": r.get("question_number", 0),
                "question": r.get("question", ""),
                "choices": r.get("choices", ""),
                "has_image": r.get("has_image", False),
            })

    if not extracted:
        continue

    # 2단계: 대분류 배정
    section_key = get_section_key(exam_type, subject)
    tax = TAXONOMY.get(section_key)
    if not tax:
        print(f"  [taxonomy 없음] {exam_type} {subject} (key={section_key})")
        continue

    large_list = list(tax.keys())
    subject_note = SUBJECT_NOTES.get(subject, "")

    large_result = {}
    for batch in chunk(extracted, BATCH_SIZE):
        res = classify_large_batch(batch, large_list, subject_note)
        large_result.update(res)
        time.sleep(1)

    # 3단계: 중분류 배정
    by_large = {}
    for it in extracted:
        large = large_result.get(it["row_key"])
        if large not in tax:
            print(f"  [경고] row_key {it['row_key']} 대분류 미확정/유효하지 않음: {large}")
            continue
        by_large.setdefault(large, []).append(it)

    mid_result = {}
    for large, sub_items in by_large.items():
        mid_list = tax[large]
        if not mid_list:
            print(f"  [알림] '{large}' 하위에 정의된 중분류 없음 — 스킵")
            continue
        for batch in chunk(sub_items, BATCH_SIZE):
            res = classify_mid_batch(batch, mid_list, subject_note)
            mid_result.update(res)
            time.sleep(1)

    # 최종 저장
    for it in extracted:
        row_key = it["row_key"]
        q_num = it["question_number"]
        large = large_result.get(row_key, "")
        mid = mid_result.get(row_key, "")
        answer = answers.get(str(q_num), "")
        row_id = f"{exam_type}_2026_2회_{subject}_{q_num}"

        if not answer:
            failed_log.append({
                "exam_type": exam_type, "subject": subject,
                "page_num": "-", "사유": f"{q_num}번 정답 매칭 실패"
            })
        if not large or not mid:
            failed_log.append({
                "exam_type": exam_type, "subject": subject,
                "page_num": "-", "사유": f"{q_num}번 분류 실패 (large={large}, mid={mid})"
            })

        new_row_data = [
            row_id, exam_type, "2026", "2", subject, q_num,
            it["question"], it["choices"], answer,
            it["has_image"], "",
            large, mid, "",
        ]

        if row_id in id_to_row:
            r = id_to_row[row_id]
            for col_idx, value in enumerate(new_row_data, start=1):
                ws.cell(row=r, column=col_idx, value=value)
        else:
            ws.append(new_row_data)
            id_to_row[row_id] = ws.max_row

    wb.save(OUTPUT_EXCEL)
    print(f"  저장완료: {exam_type} {subject}")

# 실패 로그 저장
if failed_log:
    fail_wb = openpyxl.Workbook()
    fail_ws = fail_wb.active
    fail_ws.append(["exam_type", "subject", "page_num", "사유"])
    for item in failed_log:
        fail_ws.append([item["exam_type"], item["subject"], item["page_num"], item["사유"]])
    fail_path = OUTPUT_EXCEL.replace(".xlsx", "_실패목록.xlsx")
    fail_wb.save(fail_path)
    print(f"\n실패/경고 {len(failed_log)}건 → '{fail_path}'에 저장")
else:
    print("\n실패 없이 전부 성공")

print("\n전체 완료!" if not PILOT_TARGET else "\n파일럿 완료! 결과 확인 후 PILOT_TARGET = None으로 바꿔서 전체 실행하세요.")