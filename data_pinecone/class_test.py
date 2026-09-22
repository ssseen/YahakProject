import os
from dotenv import load_dotenv
load_dotenv()

import openpyxl
from pinecone import Pinecone
from collections import defaultdict, Counter

DATA_DIR = os.environ["DATA_DIR"]
PINECONE_API_KEY = os.environ["PINECONE_API_KEY"]
TEST_EXCEL_PATH = os.path.join(DATA_DIR, "테스트세트.xlsx")

INDEX_NAME = "geondi-questions"
NAMESPACE  = "questions"
MAX_K      = 10

pc = Pinecone(api_key=PINECONE_API_KEY)
index = pc.Index(INDEX_NAME)


def load_test_questions():
    wb = openpyxl.load_workbook(TEST_EXCEL_PATH, data_only=True)
    ws = wb.active
    header = [c.value for c in ws[1]]
    col = {name: idx for idx, name in enumerate(header)}

    items = []
    for row in ws.iter_rows(min_row=2):
        row_id = row[col["id"]].value
        question = row[col["question"]].value or ""
        choices = row[col["choices"]].value or ""
        if not row_id or not question:
            continue
        items.append({
            "id": row_id,
            "subject": row[col["subject"]].value or "",
            "text": f"{question}\n{choices}".strip(),
        })
    return items


def main():
    items = load_test_questions()

    # k(1~10)별 정답 개수 누적
    correct_by_k = {k: 0 for k in range(1, MAX_K + 1)}
    total = 0

    for i, it in enumerate(items, start=1):
        response = index.search_records(
            namespace=NAMESPACE,
            query={"inputs": {"text": it["text"]}, "top_k": MAX_K},  # 검색은 한 번만, top_10까지
        )
        hits = response.result.hits
        if not hits:
            total += 1
            continue

        subjects_in_order = [h.fields.get("subject") for h in hits]

        # k=1부터 10까지, 앞에서 k개만 잘라서 다수결 계산
        for k in range(1, MAX_K + 1):
            subset = subjects_in_order[:k]
            if not subset:
                continue
            predicted = Counter(subset).most_common(1)[0][0]
            if predicted == it["subject"]:
                correct_by_k[k] += 1

        total += 1

        if i % 50 == 0:
            print(f"진행중: {i}/{len(items)}")

    print(f"\n=== top_k별 과목(subject) 일치율 (필터 없음, 다수결) ===")
    print(f"{'k':<6}{'정확도':<10}")
    for k in range(1, MAX_K + 1):
        acc = correct_by_k[k] / total * 100 if total else 0
        print(f"top_{k:<3}{acc:<9.1f}%")


if __name__ == "__main__":
    main()