import os
import re
import unicodedata
import openpyxl

# ── 설정 ──────────────────────────────────────────
EXCEL_PATH = "/Users/parkseyeon/학교/졸업프로젝트/data/유형분류.xlsx"
IMAGE_DIR = "/Users/parkseyeon/학교/졸업프로젝트/data/이미지"
GITHUB_BASE_URL = "https://raw.githubusercontent.com/ssseen/yahak_data/main"


def nfc(s):
    """한글 유니코드 정규화 (NFD → NFC) — macOS 파일명 이슈 방지"""
    return unicodedata.normalize("NFC", s) if isinstance(s, str) else s


# ── 이미지 폴더 파싱: 파일 하나가 여러 문제를 커버하는 경우까지 처리 ──
def parse_image_files(image_dir, subject):
    """subject 폴더 안 파일들을 훑어서
       {(exam_type, year, exam_round, subject, question_number): 파일명} 매핑 생성"""
    folder = os.path.join(image_dir, subject)
    mapping = {}
    if not os.path.isdir(folder):
        print(f"  [폴더 없음] {folder}")
        return mapping

    # 예: 중졸_2025_2회_국어_17-19.png  또는  고졸_2021_1회_수학_7.png
    pattern = re.compile(r"^(.+?)_(\d{4})_(\d+)회_(.+?)_(\d+)(?:-(\d+))?$")

    for fname in os.listdir(folder):
        if not fname.lower().endswith(".png"):
            continue
        name = nfc(fname[:-4])  # 확장자 제거 + 정규화

        m = pattern.match(name)
        if not m:
            print(f"  [패턴 매칭 실패] {subject}/{fname}")
            continue

        exam_type, year, exam_round, subj, start, end = m.groups()
        exam_type = nfc(exam_type)
        subj = nfc(subj)
        start = int(start)
        end = int(end) if end else start

        for qnum in range(start, end + 1):
            key = (exam_type, year, exam_round, subj, str(qnum))
            mapping[key] = nfc(fname)

    return mapping


# ── 엑셀 열기 ──────────────────────────────────────
wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb.active

header = [c.value for c in ws[1]]
col = {name: idx + 1 for idx, name in enumerate(header)}  # 1-based

required = ["id", "subject", "exam_type", "year", "exam_round", "question_number"]
for r in required:
    if r not in col:
        raise ValueError(f"엑셀에 '{r}' 컬럼이 없습니다. 헤더 확인: {header}")

if "has_image" not in col:
    col["has_image"] = ws.max_column + 1
    ws.cell(row=1, column=col["has_image"], value="has_image")
if "image_path" not in col:
    col["image_path"] = ws.max_column + 1
    ws.cell(row=1, column=col["image_path"], value="image_path")

# ── 과목별 이미지 매핑을 미리 만들어두기 (한 번씩만 파싱) ──
image_maps = {}

# ── 행별로 매칭 ──────────────────────────────────
count_found = 0
count_total = 0
for row in ws.iter_rows(min_row=2):
    row_id = row[col["id"] - 1].value
    subject = nfc(row[col["subject"] - 1].value)
    exam_type = nfc(row[col["exam_type"] - 1].value)
    year = str(row[col["year"] - 1].value or "")
    exam_round = str(row[col["exam_round"] - 1].value or "")
    question_number = str(row[col["question_number"] - 1].value or "")

    if not row_id or not subject:
        continue

    count_total += 1

    if subject not in image_maps:
        print(f"'{subject}' 폴더 파싱 중...")
        image_maps[subject] = parse_image_files(IMAGE_DIR, subject)

    key = (exam_type, year, exam_round, subject, question_number)
    r = row[0].row

    if key in image_maps[subject]:
        fname = image_maps[subject][key]
        ws.cell(row=r, column=col["has_image"], value=1)
        ws.cell(row=r, column=col["image_path"], value=f"{GITHUB_BASE_URL}/{subject}/{fname}")
        count_found += 1
    else:
        ws.cell(row=r, column=col["has_image"], value=0)
        ws.cell(row=r, column=col["image_path"], value="")

wb.save(EXCEL_PATH)
print(f"\n전체 문제 {count_total}건 중 이미지 매칭 {count_found}건")