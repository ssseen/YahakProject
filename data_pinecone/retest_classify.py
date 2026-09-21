import os
from dotenv import load_dotenv
load_dotenv()

import re
import json
import base64
import time
import openpyxl
from anthropic import Anthropic
from pdf2image import convert_from_path
from io import BytesIO

DATA_DIR = os.environ["DATA_DIR"]
ANTHROPIC_API_KEY = os.environ["ANTHROPIC_API_KEY"]
TEST_FOLDER  = os.path.join(DATA_DIR, "test_folder")
OUTPUT_EXCEL = os.path.join(DATA_DIR, "테스트세트.xlsx")

MODEL = "claude-sonnet-5"   # 상위 모델로 재추출
EXAM_TYPES = ["고졸", "중졸", "초졸"]
SUBJECT_ORDER = ["국어", "수학", "영어", "사회", "과학"]

client = Anthropic(api_key=ANTHROPIC_API_KEY)

def main():
    import shutil
    backup_path = OUTPUT_EXCEL.replace(".xlsx", "_backup.xlsx")
    shutil.copy(OUTPUT_EXCEL, backup_path)
    print(f"백업 생성: {backup_path}")

    wb = openpyxl.load_workbook(OUTPUT_EXCEL)
    ws = wb.active
    ...

def image_to_base64(img):
    buffer = BytesIO()
    img.save(buffer, format="PNG")
    return base64.b64encode(buffer.getvalue()).decode("utf-8")


def extract_page(page_image, info, page_num, max_retry=3):
    tool = {
        "name": "extract_questions",
        "description": "페이지 안의 각 문제를 정확히 추출한다.",
        "input_schema": {
            "type": "object",
            "properties": {
                "questions": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "question_number": {"type": "integer"},
                            "question": {"type": "string"},
                            "choices": {"type": "string"},
                            "has_image": {"type": "boolean"},
                        },
                        "required": ["question_number", "question", "choices", "has_image"],
                    },
                }
            },
            "required": ["questions"],
        },
    }

    prompt = f"""아래는 검정고시 {info['exam_type']} {info['subject']} 시험지 {page_num}페이지 이미지입니다.
이 페이지에 있는 각 문제의 문제 텍스트(보기 제외)와 보기를 정확하게, 원문 그대로 추출하세요.
오타나 누락 없이 이미지에 보이는 텍스트를 정확히 옮기는 것이 가장 중요합니다.
- has_image는 문제에 도형/표/그래프 등 이미지가 있으면 true
- 보기가 이미지(도형/그래프 등)인 경우 choices는 빈 문자열

반드시 extract_questions 도구를 호출하세요."""

    b64 = image_to_base64(page_image)

    for attempt in range(max_retry):
        try:
            resp = client.messages.create(
                model=MODEL,
                max_tokens=4000,
                tools=[tool],
                tool_choice={"type": "tool", "name": "extract_questions"},
                messages=[{
                    "role": "user",
                    "content": [
                        {"type": "image", "source": {"type": "base64", "media_type": "image/png", "data": b64}},
                        {"type": "text", "text": prompt}
                    ]
                }],
            )
            for block in resp.content:
                if block.type == "tool_use":
                    return block.input["questions"]
        except Exception as e:
            print(f"  [재시도 {attempt+1}/{max_retry}] {e}")
            time.sleep(10)

    return []


def main():
    wb = openpyxl.load_workbook(OUTPUT_EXCEL)
    ws = wb.active

    header = [c.value for c in ws[1]]
    col = {name: idx + 1 for idx, name in enumerate(header)}  # 1-based

    # id → 행번호 매핑
    id_to_row = {}
    for row in ws.iter_rows(min_row=2):
        row_id = row[col["id"] - 1].value
        if row_id:
            id_to_row[row_id] = row[0].row

    for exam_type in EXAM_TYPES:
        for subject in SUBJECT_ORDER:
            pdf_path = os.path.join(TEST_FOLDER, f"2026년도_2차시험_{exam_type}학력_{subject}_문제.pdf")
            if not os.path.exists(pdf_path):
                print(f"[파일없음] {pdf_path}")
                continue

            print(f"\n재추출중: {exam_type} {subject}")
            pages = convert_from_path(pdf_path)
            info = {"exam_type": exam_type, "subject": subject}

            for page_num, page_img in enumerate(pages, start=1):
                print(f"  페이지 {page_num}/{len(pages)}")
                results = extract_page(page_img, info, page_num)
                if not results:
                    print(f"  [스킵] 페이지 {page_num} 추출 실패")
                    continue
                time.sleep(3)

                for r in results:
                    q_num = r.get("question_number", 0)
                    row_id = f"{exam_type}_2026_2회_{subject}_{q_num}"

                    if row_id not in id_to_row:
                        print(f"  [경고] {row_id} 기존 엑셀에 없음 — 건너뜀")
                        continue

                    row_num = id_to_row[row_id]
                    # question, choices, has_image만 덮어씀 (category는 건드리지 않음)
                    ws.cell(row=row_num, column=col["question"], value=r.get("question", ""))
                    ws.cell(row=row_num, column=col["choices"], value=r.get("choices", ""))
                    ws.cell(row=row_num, column=col["has_image"], value=r.get("has_image", False))

            wb.save(OUTPUT_EXCEL)
            print(f"  저장완료: {exam_type} {subject}")

    print("\n전체 재추출 완료!")


if __name__ == "__main__":
    main()