import os
from dotenv import load_dotenv
load_dotenv()

import openpyxl
from pinecone import Pinecone

# ── 설정 ──────────────────────────────────────────
DATA_DIR = os.environ["DATA_DIR"]

PINECONE_API_KEY = os.environ["PINECONE_API_KEY"]
TEST_EXCEL_PATH = os.path.join(DATA_DIR, "테스트세트.xlsx")

INDEX_NAME = "geondi-questions"
NAMESPACE  = "questions"
TOP_K      = 5

pc = Pinecone(api_key=PINECONE_API_KEY)
index = pc.Index(INDEX_NAME)


def load_test_questions():
    wb = openpyxl.load_workbook(TEST_EXCEL_PATH, data_only=True)
    ws = wb.active

    header = [c.value for c in ws[1]]
    col = {name: idx for idx, name in enumerate(header)}

    required = ["id", "subject", "question", "choices", "category_large", "category_mid"]
    for r in required:
        if r not in col:
            raise ValueError(f"엑셀에 '{r}' 컬럼이 없습니다. 헤더 확인: {header}")

    items = []
    for row in ws.iter_rows(min_row=2):
        row_id = row[col["id"]].value
        question = row[col["question"]].value or ""
        choices = row[col["choices"]].value or ""
        if not row_id or not question:
            continue

        items.append({
            "id": row_id,
            "subject": row[col["subject"]].value or "",
            "text": f"{question}\n{choices}".strip(),
            "category_large": row[col["category_large"]].value or "",
            "category_mid": row[col["category_mid"]].value or "",
        })

    return items


def search_similar(text: str, top_k: int = TOP_K):
    response = index.search_records(
        namespace=NAMESPACE,
        query={"inputs": {"text": text}, "top_k": top_k},
    )
    return response.result.hits  # Hit 객체 리스트


def main():
    items = load_test_questions()
    print(f"테스트 문제 {len(items)}건 로드")

    per_question_log = []
    total_large_correct = 0
    total_mid_correct = 0
    total_hits = 0

    for i, it in enumerate(items, start=1):
        hits = search_similar(it["text"])
        if not hits:
            print(f"  [{i}/{len(items)}] {it['id']} — 검색 결과 없음")
            per_question_log.append({
                "id": it["id"], "subject": it["subject"],
                "true_large": it["category_large"], "true_mid": it["category_mid"],
                "large_acc": 0, "mid_acc": 0, "hit_count": 0,
            })
            continue

        large_correct = sum(
            1 for h in hits
            if h.fields.get("category_large") == it["category_large"]
        )
        mid_correct = sum(
            1 for h in hits
            if h.fields.get("category_mid") == it["category_mid"]
        )

        large_acc = large_correct / len(hits)
        mid_acc = mid_correct / len(hits)

        total_large_correct += large_correct
        total_mid_correct += mid_correct
        total_hits += len(hits)

        per_question_log.append({
            "id": it["id"], "subject": it["subject"],
            "true_large": it["category_large"], "true_mid": it["category_mid"],
            "large_acc": round(large_acc, 2), "mid_acc": round(mid_acc, 2),
            "hit_count": len(hits),
        })

        if i % 20 == 0:
            print(f"  진행중: {i}/{len(items)}")

    overall_large_acc = total_large_correct / total_hits if total_hits else 0
    overall_mid_acc = total_mid_correct / total_hits if total_hits else 0

    print(f"\n=== 전체 결과 (top_k={TOP_K}) ===")
    print(f"category_large 정확도: {overall_large_acc*100:.1f}%")
    print(f"category_mid   정확도: {overall_mid_acc*100:.1f}%")

    log_path = TEST_EXCEL_PATH.replace(".xlsx", "_검증결과.xlsx")
    log_wb = openpyxl.Workbook()
    log_ws = log_wb.active
    log_ws.append(["id", "subject", "true_large", "true_mid", "large_acc", "mid_acc", "hit_count"])
    for log in per_question_log:
        log_ws.append([
            log["id"], log["subject"], log["true_large"], log["true_mid"],
            log["large_acc"], log["mid_acc"], log["hit_count"],
        ])
    log_wb.save(log_path)
    print(f"\n문제별 상세 결과 → '{log_path}'에 저장")


if __name__ == "__main__":
    main()