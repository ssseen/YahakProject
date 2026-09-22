import os
from dotenv import load_dotenv
load_dotenv()

import openpyxl
from pinecone import Pinecone
from collections import defaultdict, Counter

# ── 설정 ──────────────────────────────────────────
DATA_DIR = os.environ["DATA_DIR"]

PINECONE_API_KEY = os.environ["PINECONE_API_KEY"]
TEST_EXCEL_PATH = os.path.join(DATA_DIR, "테스트세트.xlsx")

INDEX_NAME = "geondi-questions"
NAMESPACE  = "questions"
TOP_K      = 10   # 다수결에 쓸 후보 개수

pc = Pinecone(api_key=PINECONE_API_KEY)
index = pc.Index(INDEX_NAME)


def load_test_questions():
    wb = openpyxl.load_workbook(TEST_EXCEL_PATH, data_only=True)
    ws = wb.active

    header = [c.value for c in ws[1]]
    col = {name: idx for idx, name in enumerate(header)}

    required = ["id", "subject", "question", "choices", "category_large", "category_mid"]
    for r in required:
        if r not in col:
            raise ValueError(f"엑셀에 '{r}' 컬럼이 없습니다. 헤더 확인: {header}")

    items = []
    for row in ws.iter_rows(min_row=2):
        row_id = row[col["id"]].value
        question = row[col["question"]].value or ""
        choices = row[col["choices"]].value or ""
        if not row_id or not question:
            continue

        items.append({
            "id": row_id,
            "subject": row[col["subject"]].value or "",
            "text": f"{question}\n{choices}".strip(),
            "category_large": row[col["category_large"]].value or "",
            "category_mid": row[col["category_mid"]].value or "",
        })

    return items


def search_similar_filtered(text: str, subject: str, top_k: int = TOP_K):
    response = index.search_records(
        namespace=NAMESPACE,
        query={
            "inputs": {"text": text},
            "top_k": top_k,
            "filter": {"subject": {"$eq": subject}},
        },
    )
    return response.result.hits


def main():
    items = load_test_questions()
    print(f"테스트 문제 {len(items)}건 로드 (subject 필터 + 다수결 top_{TOP_K})")

    per_question_log = []

    large_correct_count = 0
    mid_correct_count = 0
    total = 0

    by_subject = defaultdict(lambda: {"count": 0, "large": 0, "mid": 0})

    for i, it in enumerate(items, start=1):
        hits = search_similar_filtered(it["text"], it["subject"])
        subj = it["subject"]

        if not hits:
            print(f"  [{i}/{len(items)}] {it['id']} — 검색 결과 없음")
            per_question_log.append({
                "id": it["id"], "subject": subj,
                "true_large": it["category_large"], "true_mid": it["category_mid"],
                "predicted_large": "", "predicted_mid": "",
                "large_correct": 0, "mid_correct": 0,
            })
            by_subject[subj]["count"] += 1
            continue

        # ── 다수결: top_k 중 가장 많이 나온 category를 예측값으로 채택 ──
        large_votes = Counter(h.fields.get("category_large") for h in hits)
        mid_votes = Counter(h.fields.get("category_mid") for h in hits)

        predicted_large = large_votes.most_common(1)[0][0]
        predicted_mid = mid_votes.most_common(1)[0][0]

        large_ok = 1 if predicted_large == it["category_large"] else 0
        mid_ok = 1 if predicted_mid == it["category_mid"] else 0

        large_correct_count += large_ok
        mid_correct_count += mid_ok
        total += 1

        by_subject[subj]["count"] += 1
        by_subject[subj]["large"] += large_ok
        by_subject[subj]["mid"] += mid_ok

        per_question_log.append({
            "id": it["id"], "subject": subj,
            "true_large": it["category_large"], "true_mid": it["category_mid"],
            "predicted_large": predicted_large, "predicted_mid": predicted_mid,
            "large_correct": large_ok, "mid_correct": mid_ok,
        })

        if i % 20 == 0:
            print(f"  진행중: {i}/{len(items)}")

    overall_large = large_correct_count / total if total else 0
    overall_mid = mid_correct_count / total if total else 0

    print(f"\n=== 전체 결과 (다수결 방식, top_{TOP_K} 투표) ===")
    print(f"category_large 정확도: {overall_large*100:.1f}%")
    print(f"category_mid   정확도: {overall_mid*100:.1f}%")

    print(f"\n=== 과목별 결과 ===")
    print(f"{'과목':<6}{'문항수':<8}{'large':<10}{'mid':<10}")
    for subj, stats in sorted(by_subject.items()):
        c = stats["count"]
        if c == 0:
            continue
        print(f"{subj:<6}{c:<8}{stats['large']/c*100:<9.1f}%{stats['mid']/c*100:<9.1f}%")

    log_path = TEST_EXCEL_PATH.replace(".xlsx", "_검증결과_다수결.xlsx")
    log_wb = openpyxl.Workbook()
    log_ws = log_wb.active
    log_ws.append(["id", "subject", "true_large", "true_mid", "predicted_large", "predicted_mid",
                   "large_correct", "mid_correct"])
    for log in per_question_log:
        log_ws.append([
            log["id"], log["subject"], log["true_large"], log["true_mid"],
            log["predicted_large"], log["predicted_mid"],
            log["large_correct"], log["mid_correct"],
        ])
    log_wb.save(log_path)
    print(f"\n문제별 상세 결과 → '{log_path}'에 저장")


if __name__ == "__main__":
    main()