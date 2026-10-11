"""
Pinecone 분류 방식 비교: 다수결 vs 가중투표 vs top-1

사용법 (.env가 있는 data_pinecone 폴더에서):
  python eval_compare.py collect   # 테스트 문제를 검색해 search_cache.json 저장 (API 호출은 여기서만)
  python eval_compare.py eval      # 캐시로 방식별 정확도 계산 (API 호출 없음, 몇 초)

collect는 중간에 끊겨도 다시 실행하면 이어서 합니다.
"""
import os
import sys
import csv
import json
import time
from collections import Counter, defaultdict

INDEX_NAME = "geondi-questions"
NAMESPACE = "questions"
TOP_K = 10                      # 캐시에 저장할 후보 개수 (k를 바꿔 비교할 수 있게 넉넉히)
CACHE_PATH = "search_cache.json"
SUMMARY_PATH = "compare_summary.csv"
ERRORS_PATH = "errors_top1.csv"


# ───────────────────────── collect ─────────────────────────
def _get(obj, key, default=None):
    try:
        return obj[key]
    except (KeyError, TypeError, AttributeError):
        return getattr(obj, key, default)


def hit_to_dict(h):
    f = _get(h, "fields", {}) or {}
    score = _get(h, "_score")
    if score is None:
        score = _get(h, "score", 0.0)
    return {
        "score": float(score or 0.0),
        "subject": f.get("subject", ""),
        "large": f.get("category_large", ""),
        "mid": f.get("category_mid", ""),
        "orig_id": f.get("original_id", ""),
        "question": (f.get("question", "") or "")[:200],
    }


def search(index, text, top_k, subject=None):
    query = {"inputs": {"text": text}, "top_k": top_k}
    if subject:
        query["filter"] = {"subject": {"$eq": subject}}
    for attempt in range(3):
        try:
            res = index.search_records(namespace=NAMESPACE, query=query)
            return [hit_to_dict(h) for h in res.result.hits]
        except Exception as e:  # 일시적 오류 재시도
            print(f"    검색 오류({attempt + 1}/3): {e}")
            time.sleep(2)
    return []


def load_test_questions():
    import openpyxl
    data_dir = os.environ["DATA_DIR"]
    path = os.path.join(data_dir, "테스트세트.xlsx")
    ws = openpyxl.load_workbook(path, data_only=True).active
    header = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(header)}
    for r in ["id", "subject", "question", "choices", "category_large", "category_mid"]:
        if r not in col:
            raise ValueError(f"엑셀에 '{r}' 컬럼이 없습니다. 헤더: {header}")
    items = []
    for row in ws.iter_rows(min_row=2):
        row_id = row[col["id"]].value
        question = row[col["question"]].value or ""
        choices = row[col["choices"]].value or ""
        if not row_id or not question:
            continue
        items.append({
            "id": str(row_id),
            "subject": row[col["subject"]].value or "",
            # 업서트와 동일한 입력 텍스트 (문제 + 보기)
            "text": f"{question}\n{choices}".strip(),
            "true_large": row[col["category_large"]].value or "",
            "true_mid": row[col["category_mid"]].value or "",
        })
    return items


def collect():
    from dotenv import load_dotenv
    load_dotenv()
    from pinecone import Pinecone

    index = Pinecone(api_key=os.environ["PINECONE_API_KEY"]).Index(INDEX_NAME)
    items = load_test_questions()

    cache = {}
    if os.path.exists(CACHE_PATH):
        with open(CACHE_PATH, encoding="utf-8") as f:
            cache = {c["id"]: c for c in json.load(f)}
        print(f"기존 캐시 {len(cache)}건 → 이어서 진행")

    print(f"테스트 문제 {len(items)}건 검색 시작 (subject 필터 O/X 각각 top_{TOP_K})")
    for i, it in enumerate(items, start=1):
        if it["id"] in cache:
            continue
        cache[it["id"]] = {
            "id": it["id"],
            "subject": it["subject"],
            "true_large": it["true_large"],
            "true_mid": it["true_mid"],
            "text": it["text"][:300],
            "filtered": search(index, it["text"], TOP_K, subject=it["subject"]),
            "nofilter": search(index, it["text"], TOP_K, subject=None),
        }
        if i % 20 == 0:
            print(f"  진행중: {i}/{len(items)}")
            with open(CACHE_PATH, "w", encoding="utf-8") as f:
                json.dump(list(cache.values()), f, ensure_ascii=False)

    with open(CACHE_PATH, "w", encoding="utf-8") as f:
        json.dump(list(cache.values()), f, ensure_ascii=False)
    print(f"완료 → {CACHE_PATH} ({len(cache)}건). 이제 `python eval_compare.py eval`")


# ───────────────────────── eval ─────────────────────────
def majority(hits, field, k):
    votes = Counter(h[field] for h in hits[:k])
    return votes.most_common(1)[0][0] if votes else ""      # 동점이면 먼저 나온(유사도 높은) 쪽


def weighted(hits, field, k):
    s = defaultdict(float)
    for h in hits[:k]:
        s[h[field]] += h["score"]
    return max(s, key=s.get) if s else ""


def top1(hits, field):
    return hits[0][field] if hits else ""


def hier(hits, field, k=5):
    """대분류는 가중투표, 중분류는 '그 대분류에 속한 후보 중 가장 유사한 것'의 중분류"""
    large = weighted(hits, "large", k)
    if field == "large":
        return large
    if field == "mid":
        for h in hits:
            if h["large"] == large:
                return h["mid"]
    return top1(hits, field)


METHODS = {
    "다수결 k=10 (현재)": lambda h, f: majority(h, f, 10),
    "다수결 k=5":          lambda h, f: majority(h, f, 5),
    "다수결 k=3":          lambda h, f: majority(h, f, 3),
    "가중투표 k=10":       lambda h, f: weighted(h, f, 10),
    "가중투표 k=5":        lambda h, f: weighted(h, f, 5),
    "계층(대분류투표→중분류)": lambda h, f: hier(h, f, 5),
    "top-1":               lambda h, f: top1(h, f),
}


def accuracy(rows, hits_key, field, truth, predictor):
    """전체 정확도 + 과목별 정확도"""
    total = ok = 0
    by_subj = defaultdict(lambda: [0, 0])
    for r in rows:
        hits = sorted(r[hits_key], key=lambda h: -h["score"])
        pred = predictor(hits, field) if hits else ""
        correct = int(pred == r[truth])
        total += 1
        ok += correct
        by_subj[r["subject"]][0] += 1
        by_subj[r["subject"]][1] += correct
    return ok / total if total else 0, {s: v[1] / v[0] for s, v in by_subj.items()}, total


def evaluate():
    with open(CACHE_PATH, encoding="utf-8") as f:
        rows = json.load(f)
    print(f"테스트 문제 {len(rows)}건\n")

    tasks = [
        ("대분류 (subject 필터 O)", "filtered", "large", "true_large"),
        ("중분류 (subject 필터 O)", "filtered", "mid", "true_mid"),
        ("과목 분류 (필터 X, 과목을 맞히는 경우)", "nofilter", "subject", "subject"),
    ]
    summary = [["task", "method", "overall"]]
    subjects = sorted({r["subject"] for r in rows})

    for title, hits_key, field, truth in tasks:
        print(f"=== {title} ===")
        print(f"{'방식':<16}{'전체':>8}" + "".join(f"{s:>8}" for s in subjects))
        for name, fn in METHODS.items():
            overall, per, _ = accuracy(rows, hits_key, field, truth, fn)
            print(f"{name:<16}{overall * 100:>7.1f}%" + "".join(f"{per.get(s, 0) * 100:>7.1f}%" for s in subjects))
            summary.append([title, name, f"{overall * 100:.2f}"])
        print()

    # 현재 방식(다수결 k=10) → top-1로 바꿨을 때 문제 몇 개가 달라지는지
    print("=== 문제 단위 비교: 다수결 k=10 → top-1 (필터 O) ===")
    for label, field, truth in (("대분류", "large", "true_large"), ("중분류", "mid", "true_mid")):
        gained = lost = 0
        for r in rows:
            hits = sorted(r["filtered"], key=lambda h: -h["score"])
            if not hits:
                continue
            a = majority(hits, field, 10) == r[truth]
            b = top1(hits, field) == r[truth]
            gained += int(b and not a)
            lost += int(a and not b)
        print(f"{label}: 맞게 된 문제 {gained}개 / 틀리게 된 문제 {lost}개 → 순증가 {gained - lost}개 (전체 {len(rows)}문제)")
    print()

    # top-1 오답 목록 + 자주 틀리는 (정답 → 예측) 대분류 쌍
    errors = [["id", "subject", "true_large", "top1_large", "true_mid", "top1_mid", "top1_score",
               "test_question", "top1_orig_id", "top1_question"]]
    pairs = Counter()
    for r in rows:
        hits = sorted(r["filtered"], key=lambda h: -h["score"])
        if not hits:
            continue
        t = hits[0]
        if t["large"] != r["true_large"] or t["mid"] != r["true_mid"]:
            errors.append([r["id"], r["subject"], r["true_large"], t["large"],
                           r["true_mid"], t["mid"], f"{t['score']:.3f}",
                           (r.get("text", "") or "").replace("\n", " "),
                           t.get("orig_id", ""),
                           (t.get("question", "") or "").replace("\n", " ")])
        if t["large"] != r["true_large"]:
            pairs[(r["subject"], r["true_large"], t["large"])] += 1
    print("=== top-1 대분류 오답 상위 10 (과목 / 정답 → 예측) ===")
    for (s, a, b), n in pairs.most_common(10):
        print(f"{n:>3}건  {s} / {a} → {b}")
    with open(ERRORS_PATH, "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows(errors)
    print(f"오답 목록 → {ERRORS_PATH}\n")

    # top-1 유사도가 낮은 문제를 걸러냈을 때 정확도 (threshold 후보 찾기용)
    print("=== top-1 유사도 하위 구간 제외 시 (대분류, 필터 O) ===")
    scored = []
    for r in rows:
        hits = sorted(r["filtered"], key=lambda h: -h["score"])
        if hits:
            scored.append((hits[0]["score"], hits[0]["large"] == r["true_large"]))
    scores = sorted(s for s, _ in scored)
    print(f"{'제외 하위%':<10}{'threshold':>11}{'남은 문항':>10}{'정확도':>9}")
    for pct in (0, 10, 25, 50):
        t = scores[int(len(scores) * pct / 100)] if scores else 0
        kept = [ok for s, ok in scored if s >= t]
        acc = sum(kept) / len(kept) if kept else 0
        print(f"{pct:<10}{t:>11.3f}{len(kept):>10}{acc * 100:>8.1f}%")

    with open(SUMMARY_PATH, "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows(summary)
    print(f"\n요약 → {SUMMARY_PATH}")


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "collect":
        collect()
    elif cmd == "eval":
        evaluate()
    else:
        print("사용법: python eval_compare.py [collect|eval]")